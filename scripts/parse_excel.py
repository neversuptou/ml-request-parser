#!/usr/bin/env python3
"""
Прогнать выгрузку заявок из админки (xlsx) через нашу модель и сохранить результат в xlsx.

Исходный текст поста берётся по ссылке:
  1) из выгрузки БД с колонкой source_text (--texts, по умолчанию applications_with_text.csv);
  2) для публичных каналов (t.me/<канал>/<id>) — со страницы Telegram-виджета (без входа в аккаунт);
  3) иначе строка помечается «нет исходного текста».

В результате рядом: наш разбор, разбор прода (DeepSeek/шаблон из выгрузки) и исходный текст.

    python scripts/parse_excel.py "D:/Download browser/zapros_za_25.09.2026.xlsx" -o out.xlsx
"""
import argparse
import csv
import html
import os
import re
import sys
import urllib.request

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from parser.extractor import _default_model, extract, render  # noqa: E402

FIELD_TITLES = [
    ("request_number", "Номер заявки"), ("project", "Наименование проекта"), ("position", "Позиция"),
    ("grade", "Уровень"), ("experience", "Опыт работы"), ("duration", "Длительность проекта"), ("rate", "Рейт"),
    ("requirements", "Требования"), ("nice_to_have", "Будет преимуществом"),
    ("tasks", "Выполняемы задачи на проекте"), ("location", "Локация проживания"), ("contacts", "Контакты"),
    ("extra", "Дополнительная информация"), ("check", "Проверить вручную"),
]
PROD_COLS = ["Номер заявки", "Наименование проекта", "Позиция", "Уровень", "Опыт работы", "Длительность проекта",
             "Рейт", "Требования", "Выполняемы задачи на проекте", "Локация проживания", "Контакты"]
_PUBLIC = re.compile(r"^https?://t\.me/([A-Za-z]\w{3,})/(\d+)$")


def fetch_public(link: str) -> str:
    """Текст поста публичного канала со страницы виджета t.me (без авторизации)."""
    m = _PUBLIC.match(link.strip())
    if not m:
        return ""
    url = f"https://t.me/{m.group(1)}/{m.group(2)}?embed=1&mode=tme"
    try:
        page = urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"}),
                                      timeout=20).read().decode("utf-8", "replace")
    except Exception:
        return ""
    t = re.search(r'class="tgme_widget_message_text[^"]*"[^>]*>(.*?)</div>', page, re.S)
    if not t:
        return ""
    body = re.sub(r"<br\s*/?>", "\n", t.group(1))
    return html.unescape(re.sub(r"<[^>]+>", "", body)).strip()


def as_text(v) -> str:
    if isinstance(v, list):
        return "\n".join(v)
    return v or ""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("-o", "--out", default="parsed.xlsx")
    ap.add_argument("--texts", default="applications_with_text.csv", help="выгрузка БД с source_text")
    args = ap.parse_args()

    texts = {}
    if os.path.exists(args.texts):
        csv.field_size_limit(10**9)
        for r in csv.DictReader(open(args.texts, encoding="utf-8")):
            if r.get("source_text"):
                texts[r["Ссылка в Telegram"].strip()] = r["source_text"]

    ws_in = openpyxl.load_workbook(args.xlsx, read_only=True).active
    rows = list(ws_in.iter_rows(values_only=True))
    head = [str(h) for h in rows[0]]
    col = {h: i for i, h in enumerate(head)}

    model = _default_model()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Разбор моделью"
    header = (["Ссылка", "Канал", "Откуда текст"] + [t for _, t in FIELD_TITLES]
              + [f"ПРОД: {c}" for c in PROD_COLS] + ["Прод обработал", "Исходный текст", "Итоговый шаблон"])
    ws.append(header)

    stats = {"db": 0, "public": 0, "none": 0}
    for r in rows[1:]:
        get = lambda name: r[col[name]] if name in col and r[col[name]] is not None else ""  # noqa: E731
        link = str(get("Ссылка в Telegram")).strip()
        text, src = texts.get(link, ""), "выгрузка БД"
        if text:
            stats["db"] += 1
        else:
            text, src = fetch_public(link), "публичный канал"
            stats["public" if text else "none"] += 1
        if not text:
            src = "нет исходного текста (приватный канал / нет ссылки)"
        fields = extract(text, model) if text else {}
        ws.append([link, get("template__title"), src]
                  + [as_text(fields.get(k)) for k, _ in FIELD_TITLES]
                  + [str(get(c)) for c in PROD_COLS] + [get("Обработано") or "шаблон", text,
                                                        render(fields) if text else ""])

    # оформление: перенос строк, ширина, выделение
    bold, grey = Font(bold=True), PatternFill("solid", fgColor="EEEEEE")
    warn = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font, c.fill = bold, grey
    check_col = header.index("Проверить вручную") + 1
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if row[check_col - 1].value:
            row[check_col - 1].fill = warn
    for i, h in enumerate(header, 1):
        width = 60 if h in ("Требования", "Выполняемы задачи на проекте", "Исходный текст", "Итоговый шаблон") else 24
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    ws.freeze_panes = "D2"
    wb.save(args.out)
    print(f"Заявок: {len(rows) - 1}; текст из БД: {stats['db']}, с публичных каналов: {stats['public']}, "
          f"без текста: {stats['none']}. Сохранено: {args.out}")


if __name__ == "__main__":
    main()
